#!/usr/bin/env python3
"""Export balance spreadsheets or CSV sheets into Godot runtime JSON."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any


EXPORT_SHEETS = {
    "Config": ("config", "key"),
    "Classes": ("classes", "class_id"),
    "Skills": ("skills", "skill_id"),
    "SkillEffects": ("skill_effects", ""),
    "StatusEffects": ("status_effects", "status_id"),
    "Enemies": ("enemies", "enemy_id"),
    "EnemyScaling": ("enemy_scaling", "scaling_id"),
    "Synergies": ("synergies", "synergy_id"),
    "SynergyEffects": ("synergy_effects", ""),
    "Waves": ("waves", "wave_id"),
}

LIST_FIELDS = {
    "stat_tags",
    "alignment_tags",
    "skill_ids",
    "skills",
    "tags",
    "thresholds",
}

INT_FIELDS = {
    "deploy_cost",
    "hp",
    "atk",
    "def",
    "dp_cost",
    "max_stack",
    "count",
    "floor",
    "wave_index",
    "effect_order",
    "stack",
    "min_tier",
    "order",
    "shop_level",
    "base_cost",
}

FLOAT_FIELDS = {
    "attack_speed",
    "move_speed",
    "attack_range",
    "search_range",
    "cooldown",
    "damage_multiplier",
    "aoe_radius",
    "heal_amount",
    "buff_defense",
    "duration",
    "tick_interval",
    "tick_value",
    "modifier_value",
    "value",
    "radius",
    "base",
    "per_layer",
    "max_value",
    "hp_multiplier",
    "atk_multiplier",
    "def_multiplier",
    "reward_multiplier",
    "delay",
}

ID_RE = re.compile(r"^[a-z0-9_]+$")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Path to .xlsx workbook or a directory of sheet CSV files.")
    parser.add_argument("--output", required=True, help="Directory for generated JSON files, usually game/data.")
    parser.add_argument("--report", default="", help="Directory for validation_report.md.")
    args = parser.parse_args()

    input_path = Path(args.input)
    output_dir = Path(args.output)
    report_dir = Path(args.report) if args.report else None

    try:
        sheets = load_input(input_path)
    except Exception as exc:
        print(f"export failed: {exc}", file=sys.stderr)
        return 2

    data, errors, warnings = build_runtime_data(sheets)
    if errors:
        write_report(report_dir, errors, warnings)
        for error in errors:
            print(error, file=sys.stderr)
        return 1

    output_dir.mkdir(parents=True, exist_ok=True)
    for name, payload in data.items():
        (output_dir / f"{name}.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent="\t") + "\n",
            encoding="utf-8",
        )

    write_report(report_dir, errors, warnings)
    print(f"exported {len(data)} json files to {output_dir}")
    return 0


def load_input(path: Path) -> dict[str, list[dict[str, Any]]]:
    if path.is_dir():
        return load_csv_directory(path)
    if path.suffix.lower() == ".xlsx":
        return load_xlsx(path)
    raise ValueError(f"unsupported input: {path}")


def load_csv_directory(path: Path) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {}
    for sheet_name in EXPORT_SHEETS:
        csv_path = path / f"{sheet_name}.csv"
        if csv_path.exists():
            result[sheet_name] = read_csv(csv_path)
    return result


def read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [normalize_row(row) for row in csv.DictReader(handle)]


def load_xlsx(path: Path) -> dict[str, list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for .xlsx export") from exc

    workbook = load_workbook(path, data_only=True)
    result: dict[str, list[dict[str, Any]]] = {}
    for sheet_name in EXPORT_SHEETS:
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            result[sheet_name] = []
            continue
        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        sheet_rows: list[dict[str, Any]] = []
        for row_values in rows[1:]:
            row = {
                headers[index]: row_values[index] if index < len(row_values) else ""
                for index in range(len(headers))
                if headers[index]
            }
            if any(value not in (None, "") for value in row.values()):
                sheet_rows.append(normalize_row(row))
        result[sheet_name] = sheet_rows
    return result


def normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    clean: dict[str, Any] = {}
    for key, value in row.items():
        field = str(key).strip()
        if not field:
            continue
        clean[field] = coerce_value(field, value)
    return clean


def coerce_value(field: str, value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        value = value.strip()
    if field == "enabled":
        return parse_bool(value, True)
    if field in LIST_FIELDS:
        if isinstance(value, list):
            return value
        return [part.strip() for part in str(value).split(",") if part.strip()]
    if field in INT_FIELDS and str(value) != "":
        return int(float(value))
    if field in FLOAT_FIELDS and str(value) != "":
        return float(value)
    return value


def parse_bool(value: Any, default: bool) -> bool:
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"true", "1", "yes", "y"}:
        return True
    if text in {"false", "0", "no", "n"}:
        return False
    return default


def build_runtime_data(sheets: dict[str, list[dict[str, Any]]]) -> tuple[dict[str, Any], list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    tables: dict[str, Any] = {}

    tables["config"] = build_config(sheets.get("Config", []))
    for sheet_name, (runtime_name, id_field) in EXPORT_SHEETS.items():
        if sheet_name in {"Config", "SkillEffects", "SynergyEffects"}:
            continue
        rows = [row for row in sheets.get(sheet_name, []) if parse_bool(row.get("enabled", True), True)]
        tables[runtime_name] = index_rows(sheet_name, rows, id_field, errors)

    merge_skill_effects(tables["skills"], sheets.get("SkillEffects", []), errors)
    merge_synergy_effects(tables["synergies"], sheets.get("SynergyEffects", []), errors)
    validate_references(tables, errors, warnings)
    validate_ranges(tables, errors, warnings)
    return tables, errors, warnings


def build_config(rows: list[dict[str, Any]]) -> dict[str, Any]:
    config: dict[str, Any] = {}
    for row in rows:
        if not parse_bool(row.get("enabled", True), True):
            continue
        key = str(row.get("key", "")).strip()
        if not key:
            continue
        value = row.get("value", "")
        value_type = str(row.get("type", "")).strip().lower()
        if value_type == "int":
            value = int(float(value))
        elif value_type == "float":
            value = float(value)
        elif value_type == "bool":
            value = parse_bool(value, False)
        config[key] = value
    return config


def index_rows(sheet_name: str, rows: list[dict[str, Any]], id_field: str, errors: list[str]) -> dict[str, dict[str, Any]]:
    table: dict[str, dict[str, Any]] = {}
    for row_index, row in enumerate(rows, start=2):
        row_id = str(row.get(id_field, "")).strip()
        if not row_id:
            errors.append(f"{sheet_name} row {row_index}: missing {id_field}")
            continue
        if not ID_RE.match(row_id):
            errors.append(f"{sheet_name} row {row_index}: invalid id {row_id}")
        if row_id in table:
            errors.append(f"{sheet_name} row {row_index}: duplicate id {row_id}")
            continue
        row[id_field] = row_id
        table[row_id] = row
    return table


def merge_skill_effects(skills: dict[str, dict[str, Any]], rows: list[dict[str, Any]], errors: list[str]) -> None:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row_index, row in enumerate(rows, start=2):
        if not parse_bool(row.get("enabled", True), True):
            continue
        skill_id = str(row.get("skill_id", "")).strip()
        if skill_id not in skills:
            errors.append(f"SkillEffects row {row_index}: missing skill reference {skill_id}")
            continue
        grouped[skill_id].append(row)
    for skill_id, effects in grouped.items():
        effects.sort(key=lambda item: int(item.get("effect_order", 0)))
        skills[skill_id]["effects"] = effects


def merge_synergy_effects(synergies: dict[str, dict[str, Any]], rows: list[dict[str, Any]], errors: list[str]) -> None:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row_index, row in enumerate(rows, start=2):
        if not parse_bool(row.get("enabled", True), True):
            continue
        synergy_id = str(row.get("synergy_id", "")).strip()
        if synergy_id not in synergies:
            errors.append(f"SynergyEffects row {row_index}: missing synergy reference {synergy_id}")
            continue
        grouped[synergy_id].append(row)
    for synergy_id, effects in grouped.items():
        synergies[synergy_id]["scalable_modifiers"] = effects


def validate_references(tables: dict[str, Any], errors: list[str], warnings: list[str]) -> None:
    skills = set(tables.get("skills", {}))
    statuses = set(tables.get("status_effects", {}))
    enemies = set(tables.get("enemies", {}))
    scaling = set(tables.get("enemy_scaling", {}))

    for class_id, row in tables.get("classes", {}).items():
        for skill_id in row.get("skill_ids", []):
            if skill_id not in skills:
                errors.append(f"Classes {class_id}: unknown skill_id {skill_id}")

    for skill_id, row in tables.get("skills", {}).items():
        for effect in row.get("effects", []):
            status_id = str(effect.get("status_id", "")).strip()
            if status_id and status_id not in statuses:
                errors.append(f"Skills {skill_id}: unknown status_id {status_id}")

    for wave_id, row in tables.get("waves", {}).items():
        enemy_id = str(row.get("enemy_id", "")).strip()
        scaling_id = str(row.get("scaling_id", "")).strip()
        if enemy_id and enemy_id not in enemies:
            errors.append(f"Waves {wave_id}: unknown enemy_id {enemy_id}")
        if scaling_id and scaling_id not in scaling:
            errors.append(f"Waves {wave_id}: unknown scaling_id {scaling_id}")

    if not tables.get("classes"):
        warnings.append("Classes table is empty.")
    if not tables.get("skills"):
        warnings.append("Skills table is empty.")


def validate_ranges(tables: dict[str, Any], errors: list[str], warnings: list[str]) -> None:
    for table_name in ("classes", "enemies"):
        for row_id, row in tables.get(table_name, {}).items():
            require_positive(table_name, row_id, row, "hp", errors)
            require_non_negative(table_name, row_id, row, "atk", errors)
            require_non_negative(table_name, row_id, row, "def", errors)
            require_positive(table_name, row_id, row, "attack_speed", errors)
    for skill_id, row in tables.get("skills", {}).items():
        require_non_negative("skills", skill_id, row, "cooldown", errors)
        require_non_negative("skills", skill_id, row, "dp_cost", errors)
    for status_id, row in tables.get("status_effects", {}).items():
        if str(row.get("status_type", "")) in {"dot", "hot"} and float(row.get("tick_interval", 0.0)) <= 0.0:
            errors.append(f"status_effects {status_id}: tick_interval must be > 0 for DOT/HOT")
    if warnings:
        return


def require_positive(table: str, row_id: str, row: dict[str, Any], field: str, errors: list[str]) -> None:
    if field in row and float(row[field]) <= 0:
        errors.append(f"{table} {row_id}: {field} must be > 0")


def require_non_negative(table: str, row_id: str, row: dict[str, Any], field: str, errors: list[str]) -> None:
    if field in row and float(row[field]) < 0:
        errors.append(f"{table} {row_id}: {field} must be >= 0")


def write_report(report_dir: Path | None, errors: list[str], warnings: list[str]) -> None:
    if not report_dir:
        return
    report_dir.mkdir(parents=True, exist_ok=True)
    lines = ["# Balance Validation Report", ""]
    lines.append("## Errors")
    lines.extend(f"- {item}" for item in errors) if errors else lines.append("- None")
    lines.extend(["", "## Warnings"])
    lines.extend(f"- {item}" for item in warnings) if warnings else lines.append("- None")
    (report_dir / "validation_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    raise SystemExit(main())
